"""Synthetic local samples; never includes personal source documents."""
import io
import sys
import zipfile
from pathlib import Path
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

root = Path(sys.argv[1])
root.mkdir(parents=True, exist_ok=True)
doc = Document()
doc.add_heading("Document heading", 0)
doc.add_paragraph("A little bird sings. 😀")
table = doc.add_table(rows=1, cols=2)
table.cell(0, 0).text, table.cell(0, 1).text = "Instrument", "Bamboo flute"
doc.save(root / "notes.docx")
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2)).text = "Slide one: birds and flutes"
prs.save(root / "slides.pptx")
book = Workbook()
book.active.title = "乐器"
book.active.append(["名称", "数量"])
book.active.append(["竹笛", 3])
book.active["C2"] = "=B2*2"
book.create_sheet("隐藏").sheet_state = "hidden"
book.save(root / "table.xlsx")
pdf = PdfWriter()
page = pdf.add_blank_page(width=612, height=792)
font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): pdf._add_object(font)})})
stream = DecodedStreamObject()
stream.set_data(b"BT /F1 12 Tf 50 700 Td (PDF page one: bamboo flute.) Tj ET")
page[NameObject("/Contents")] = pdf._add_object(stream)
pdf.add_blank_page(width=612, height=792)
pdf.write(root / "mixed.pdf")
pdf2 = PdfWriter()
pdf2.add_blank_page(width=612, height=792)
pdf2.write(root / "scan.pdf")
pdf3 = PdfWriter()
pdf3.add_blank_page(width=612, height=792)
pdf3.encrypt("private-password")
pdf3.write(root / "encrypted.pdf")
(root / "multi.csv").write_text('标题,内容\n鸟,"第一行\n第二行 😀"\n', encoding="utf-8")
(root / "large.csv").write_text('"' + 'A' * 1100000 + '"\n', encoding="utf-8")
with zipfile.ZipFile(root / "book.epub", "w") as z:
    z.writestr("META-INF/container.xml", '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="EPUB/book.opf"/></rootfiles></container>')
    z.writestr("EPUB/book.opf", '<package xmlns="http://www.idpf.org/2007/opf"><manifest><item id="b" href="two.xhtml"/><item id="a" href="one.xhtml"/></manifest><spine><itemref idref="a"/><itemref idref="b"/></spine></package>')
    z.writestr("EPUB/one.xhtml", '<html><body><h1>First chapter</h1><p>Bird songs.</p><script>DO NOT KEEP</script></body></html>')
    z.writestr("EPUB/two.xhtml", '<html><body><h1>Second chapter</h1><p>Bamboo flutes.</p></body></html>')
with zipfile.ZipFile(root / "unsafe.epub", "w") as z:
    z.writestr("../outside.txt", "not allowed")
with zipfile.ZipFile(root / "entities.docx", "w") as z:
    z.writestr("word/document.xml", '<!DOCTYPE x [<!ENTITY x SYSTEM "file:///etc/passwd">]><x>&x;</x>')
print("generated synthetic fixtures")
