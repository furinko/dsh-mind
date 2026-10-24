"""Bounded local document worker. Only the private input copy is readable input.

stdout is a JSON protocol. Never fetch links, extract archives onto disk, execute
Office macros/formulas, or invoke a model. Locations refer to the original file.
"""
import csv
import importlib.util
import io
import json
import posixpath
import re
import sys
import warnings
import zipfile
from html.parser import HTMLParser
from pathlib import PurePosixPath
from xml.etree import ElementTree as ET

VERSION = "mind-documents-1"
MODULES = {".pdf": "pypdf", ".docx": "docx", ".pptx": "pptx", ".xlsx": "openpyxl", ".csv": None, ".tsv": None, ".epub": None}


class ParserError(Exception):
    def __init__(self, code, message):
        super().__init__(message)
        self.code = code


class TextHTML(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts, self.skip = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "noscript"):
            self.skip += 1
        if tag in ("p", "div", "br", "h1", "h2", "h3", "li", "tr"):
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style", "noscript"):
            self.skip = max(0, self.skip - 1)

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def xml(data):
    if b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
        raise ParserError("INFORMATION_ARCHIVE_REFUSED", "文档包含不支持的 XML 实体声明")
    return ET.fromstring(data)


def archive(data):
    z = zipfile.ZipFile(io.BytesIO(data))
    infos = z.infolist()
    if len(infos) > 10000 or sum(i.file_size for i in infos) > 128 * 1024 * 1024:
        z.close()
        raise ParserError("INFORMATION_ARCHIVE_LIMIT", "文档解压后超过128MiB或10000项")
    for item in infos:
        p = PurePosixPath(item.filename)
        if p.is_absolute() or ".." in p.parts or "\\" in item.filename or item.flag_bits & 1:
            z.close()
            raise ParserError("INFORMATION_ARCHIVE_REFUSED", "文档包含不支持的归档路径或加密项")
        if item.file_size > 32 * 1024 * 1024:
            z.close()
            raise ParserError("INFORMATION_ARCHIVE_LIMIT", "文档内部单项超过32MiB")
        if item.filename.lower().endswith((".xml", ".rels")):
            # Validate before third-party XML parsers inspect the container.
            xml(z.read(item))
    return z


class Output:
    def __init__(self, request):
        self.limit = min(int(request.get("maxChars", 2000000)), 2000000)
        self.parts, self.locations, self.warnings = [], [], []
        self.length, self.partial = 0, False

    def warn(self, message):
        self.partial = True
        if len(self.warnings) < 20:
            self.warnings.append(message)

    def add(self, label, text):
        text = str(text).strip()
        if not text:
            return
        if len(self.locations) >= 10000 or self.length >= self.limit:
            self.warn("达到解析输出上限，后续内容未解析")
            raise StopIteration
        part = "\n\n[" + label + "]\n" + text
        encoded = part.encode("utf-16-le")
        if self.length + len(encoded) // 2 > self.limit:
            part = encoded[:(self.limit - self.length) * 2].decode("utf-16-le", errors="ignore")
            self.warn("达到解析字符上限，最后一个位置仅保留片段")
        start = self.length
        self.parts.append(part)
        self.length += len(part.encode("utf-16-le")) // 2
        self.locations.append({"label": label[:240], "start": start, "end": self.length})


def extract(request):
    ext = request["extension"]
    if ext not in MODULES:
        raise ParserError("INFORMATION_FORMAT_UNSUPPORTED", "没有此格式的解析器")
    module = MODULES[ext]
    if module and not importlib.util.find_spec(module):
        raise ParserError("INFORMATION_PARSER_MISSING", "缺少本地文档解析组件，请安装解析组件后重试")
    with open(request["path"], "rb") as f:
        data = f.read(50 * 1024 * 1024 + 1)
    if len(data) > 50 * 1024 * 1024:
        raise ParserError("INFORMATION_FILE_LIMIT", "原文件超过50MiB")
    if ext == ".pdf" and not data.startswith(b"%PDF-"):
        raise ParserError("INFORMATION_FORMAT_INVALID", "文件内容不是 PDF")
    if ext in (".docx", ".pptx", ".xlsx", ".epub"):
        with archive(data):
            pass
    out = Output(request)
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ParserError("INFORMATION_ENCRYPTED", "加密 PDF 需由本人提供未加密副本")
            for index, page in enumerate(reader.pages):
                if index >= 1000:
                    out.warn("仅处理前1000页")
                    break
                text = page.extract_text() or ""
                if text.strip():
                    out.add("第%d页" % (index + 1), text)
                else:
                    out.warn("第%d页无可提取文字，需 OCR 或视觉解析" % (index + 1))
            out.warnings.append("仅提取 PDF 文字层，图表、图片与版面关系未作视觉理解")
        elif ext == ".docx":
            from docx import Document
            from docx.table import Table
            doc = Document(io.BytesIO(data))
            for index, block in enumerate(doc.iter_inner_content()):
                if isinstance(block, Table):
                    for row_index, row in enumerate(block.rows):
                        out.add("正文块%d 表格行%d" % (index + 1, row_index + 1), " | ".join("列%d: %s" % (i + 1, c.text) for i, c in enumerate(row.cells)))
                else:
                    out.add("正文块%d" % (index + 1), block.text)
            out.warnings.append("仅正文文字和表格；图片、批注、修订及页眉页脚未覆盖，正文块不是页码")
        elif ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(io.BytesIO(data))
            for index, slide in enumerate(prs.slides):
                for n, shape in enumerate(slide.shapes):
                    label = "幻灯片%d 元素%d" % (index + 1, n + 1)
                    if shape.has_text_frame:
                        out.add(label, shape.text)
                    if shape.has_table:
                        for r, row in enumerate(shape.table.rows):
                            out.add(label + " 表格行%d" % (r + 1), " | ".join(c.text for c in row.cells))
            out.warnings.append("仅幻灯片文字和表格；图片、图形含义、动画和讲者备注未覆盖")
        elif ext == ".xlsx":
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            try:
                for sheet in book:
                    if sheet.sheet_state != "visible":
                        out.warn("隐藏工作表未读取：" + sheet.title[:80])
                        continue
                    if sheet.max_row > 20000 or sheet.max_column > 256:
                        out.warn("工作表超出20000行或256列，仅解析范围内内容：" + sheet.title[:80])
                    for row in sheet.iter_rows(max_row=min(sheet.max_row, 20000), max_col=min(sheet.max_column, 256)):
                        cells = ["%s: %s" % (c.coordinate, str(c.value)) for c in row if c.value is not None]
                        if cells:
                            out.add("工作表 %s 行%d" % (sheet.title, row[0].row), " | ".join(cells))
            finally:
                book.close()
            out.warnings.append("单元格原值；公式保留表达式、不计算；样式、图表与嵌入对象未覆盖")
        elif ext in (".csv", ".tsv"):
            text = data.decode("utf-8-sig", errors="strict")
            csv.field_size_limit(2000000)
            for index, row in enumerate(csv.reader(io.StringIO(text, newline=""), delimiter="\t" if ext == ".tsv" else ",")):
                out.add("记录%d" % (index + 1), " | ".join("列%d: %s" % (i + 1, value) for i, value in enumerate(row)))
        elif ext == ".epub":
            with archive(data) as z:
                ns = {"c": "urn:oasis:names:tc:opendocument:xmlns:container", "o": "http://www.idpf.org/2007/opf"}
                rootfile = xml(z.read("META-INF/container.xml")).find(".//c:rootfile", ns).attrib["full-path"]
                opf = xml(z.read(rootfile))
                manifest = {item.attrib["id"]: item.attrib.get("href", "") for item in opf.findall("o:manifest/o:item", ns)}
                for index, item in enumerate(opf.findall("o:spine/o:itemref", ns)):
                    href = manifest[item.attrib["idref"]]
                    from urllib.parse import unquote
                    target = posixpath.normpath(str(PurePosixPath(rootfile).parent / unquote(href.split("#")[0])))
                    if target == ".." or target.startswith("../") or ":" in target or target.startswith("/"):
                        raise ParserError("INFORMATION_ARCHIVE_REFUSED", "EPUB章节引用越界")
                    parser = TextHTML()
                    parser.feed(z.read(target).decode("utf-8-sig", errors="strict"))
                    out.add("章节%d %s" % (index + 1, href[:160]), "".join(parser.parts))
            out.warnings.append("按书籍章节顺序读取文字；图片和嵌入媒体未覆盖")
    except StopIteration:
        pass
    text = "".join(out.parts)
    if not text.strip():
        raise ParserError("INFORMATION_NO_TEXT", "未取得可读正文；扫描件需 OCR，图片或媒体需对应解析器")
    return {"ok": True, "text": text, "locations": out.locations, "extraction": {"parser": VERSION, "status": "partial" if out.partial else "complete", "warnings": out.warnings[:20], "units": len(out.locations)}}


def main():
    if "--capabilities" in sys.argv:
        return {"ok": True, "parser": VERSION, "formats": {ext: module is None or importlib.util.find_spec(module) is not None for ext, module in MODULES.items()}}
    return extract(json.loads(sys.stdin.buffer.read(16000)))


if __name__ == "__main__":
    warnings.filterwarnings("ignore")
    try:
        import resource
        resource.setrlimit(resource.RLIMIT_AS, (1536 * 1024 * 1024, 1536 * 1024 * 1024))
        resource.setrlimit(resource.RLIMIT_CPU, (30, 30))
    except (ImportError, ValueError, OSError):
        pass
    try:
        result = main()
    except ParserError as error:
        result = {"ok": False, "code": error.code, "message": str(error)}
    except Exception:
        # Third-party exceptions may include source text and paths: never echo them.
        result = {"ok": False, "code": "INFORMATION_PARSE_FAILED", "message": "本地文档解析失败，文件可能损坏、格式不匹配或不受支持"}
    print(json.dumps(result, ensure_ascii=True))
